# ai_insights.py
import os
import time
import tempfile
import zipfile
import pandas as pd
from PIL import Image
from openpyxl import load_workbook
from pptx import Presentation

from google import genai
from google.genai import types

# ===================== CONFIG =====================
MODEL_NAME = "gemini-2.5-flash"
OUTPUT_FILE = "security_analysis_output.xlsx"

SYSTEM_PROMPT = (
    "You are a highly skilled security system analyst. Review the provided content. "
    "Your response MUST be in two distinct parts, separated by a unique delimiter: '---DELIMITER---'. "
    "1. **File Description**: A single, concise sentence (maximum 1 line) describing exactly what is visible or contained in the file. "
    "2. **Key Findings**: A two-line, narrative-style summary detailing the core security features and functions of the asset or content shown. "
    "DO NOT mention any vulnerabilities, risks, flaws, or weaknesses. Focus ONLY on positive or neutral descriptive characteristics and operational functions. "
    "Example Response:\n"
    "A wall-mounted electronic biometric device with fingerprint scanning, keypad, and display screen showing time and status.\n"
    "---DELIMITER---\n"
    "This system provides high-assurance access control using non-transferable biometric authentication.\n"
    "It automatically logs time-stamped entries and attempts, ensuring full operational auditability."
)
# ==================================================


def get_file_type(ext: str) -> str:
    """Return raw extension like '.png', '.xlsx', etc."""
    return ext


def extract_xlsx_content(file_path: str) -> str:
    """Extract text from all sheets (first ~50 rows per sheet) of an XLSX."""
    content = f"XLSX File Content: {os.path.basename(file_path)}\n"
    try:
        workbook = load_workbook(file_path)
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            content += f"\n--- Sheet: {sheet_name} ---\n"
            row_count = 0
            for row in sheet.iter_rows():
                if row_count >= 50:
                    content += f"... Truncated after {row_count} rows ..."
                    break
                row_data = [cell.value for cell in row]
                if any(row_data):
                    content += " | ".join([str(d) for d in row_data if d is not None]) + "\n"
                    row_count += 1
        return content
    except Exception as e:
        return f"Error reading XLSX file: {e}"


def extract_pptx_content(file_path: str) -> str:
    """Extract visible text from all slides of a PPTX."""
    content = f"PPTX Presentation Content: {os.path.basename(file_path)}\n"
    try:
        prs = Presentation(file_path)
        for i, slide in enumerate(prs.slides):
            content += f"\n--- Slide {i + 1} ---\n"
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    content += shape.text + "\n"
        return content
    except Exception as e:
        return f"Error reading PPTX file: {e}"


def initialize_client():
    """Initialize Gemini client from GEMINI_API_KEY env var."""
    api_key = os.getenv("GEMINI_API_KEY")
    if not api_key:
        raise RuntimeError("GEMINI_API_KEY environment variable not set.")
    client = genai.Client(api_key=api_key)
    return client


def analyze_file(client, file_path: str, filename: str) -> dict | None:
    """Analyze a single file and return a dict of results or None if unsupported."""
    uploaded_file = None
    file_part = None
    ext = os.path.splitext(filename)[1].lower()
    file_type = get_file_type(ext)

    results = {
        "File Name": filename,
        "File Type": file_type,
        "File Description": "Processing failed.",
        "Key Findings": "Error during processing.",
    }

    try:
        if ext in (".jpg", ".jpeg", ".png"):
            # Images: pass as PIL Image
            file_part = Image.open(file_path)

        elif ext == ".xlsx":
            # Spreadsheets: extract text
            file_part = extract_xlsx_content(file_path)

        elif ext == ".pptx":
            # Presentations: extract text
            file_part = extract_pptx_content(file_path)

        elif ext == ".pdf":
            # PDFs: upload file part to Gemini Files API
            uploaded_file = client.files.upload(file=file_path)
            file_part = uploaded_file
            time.sleep(1)  # small delay is helpful

        else:
            return None  # unsupported

        # Prepare prompt + content
        contents = [file_part, SYSTEM_PROMPT]

        # Generate content
        response = client.models.generate_content(
            model=MODEL_NAME,
            contents=contents
        )

        # Parse two-part response
        response_text = response.text.strip()
        parts = response_text.split("---DELIMITER---", 1)

        if len(parts) == 2:
            description = parts[0].strip().replace("\r", "").replace("\n", " ")
            findings = parts[1].strip()
            results["File Description"] = description
            results["Key Findings"] = findings
        else:
            results["File Description"] = (
                "Error: Could not parse response correctly. Raw response follows."
            )
            results["Key Findings"] = response_text

    except genai.errors.APIError as e:
        results["File Description"] = f"API Error: {e.args[0]}"
        results["Key Findings"] = "See File Description for error details."
    except Exception as e:
        results["File Description"] = f"Local Script Error: {e}"
        results["Key Findings"] = "See File Description for error details."
    finally:
        if uploaded_file:
            try:
                time.sleep(1)
                client.files.delete(name=uploaded_file.name)
            except Exception:
                pass

    return results


def analyze_directory_to_excel(input_dir: str, output_excel_path: str | None = None) -> str:
    """
    Recursively walk a directory, analyze supported files, and write Excel.
    Returns the Excel path. Raises RuntimeError if nothing processed.
    """
    client = initialize_client()

    SUPPORTED = {".pdf", ".jpg", ".jpeg", ".png", ".xlsx", ".pptx"}
    candidate_paths = []
    for root, _, files in os.walk(input_dir):
        for name in files:
            ext = os.path.splitext(name)[1].lower()
            if ext in SUPPORTED:
                candidate_paths.append(os.path.join(root, name))

    all_results = []
    for file_path in sorted(candidate_paths):
        filename = os.path.basename(file_path)
        result = analyze_file(client, file_path, filename)
        if result:
            all_results.append(result)
        time.sleep(1.5)  # polite pacing for API

    if not all_results:
        raise RuntimeError("No supported files processed; nothing to write.")

    df = pd.DataFrame(all_results)[
        ["File Name", "File Type", "File Description", "Key Findings"]
    ]

    output_excel_path = output_excel_path or os.path.join(input_dir, OUTPUT_FILE)
    writer = pd.ExcelWriter(output_excel_path, engine="xlsxwriter")
    df.to_excel(writer, sheet_name="Security_Analysis", index=False)
    workbook = writer.book
    worksheet = writer.sheets["Security_Analysis"]

    wrap_fmt = workbook.add_format({"text_wrap": True, "align": "top"})
    worksheet.set_column("A:A", 25)
    worksheet.set_column("B:B", 12)
    worksheet.set_column("C:C", 40, wrap_fmt)
    worksheet.set_column("D:D", 60, wrap_fmt)
    writer.close()
    return output_excel_path


def extract_zip_to_tempdir(zip_path: str) -> str:
    """Unzip into a temp directory and return that directory path."""
    tmpdir = tempfile.mkdtemp(prefix="insights_")
    with zipfile.ZipFile(zip_path, "r") as zf:
        zf.extractall(tmpdir)
    return tmpdir
