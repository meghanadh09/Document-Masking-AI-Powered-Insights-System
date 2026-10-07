"""
sanitize_pptx.py
Sanitize PowerPoint files by masking sensitive information in text, tables, and images.
Uses spaCy NER, regex patterns, and optional face detection.
"""

import re
import io
import spacy
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from PIL import Image, ImageFilter
import cv2
import numpy as np

# Load spaCy model
try:
    nlp = spacy.load("en_core_web_sm")
except OSError:
    raise RuntimeError(
        "spaCy model 'en_core_web_sm' not found. Install with:\n"
        "  python -m spacy download en_core_web_sm"
    )

# Configuration
MASK_FACES = True
BLUR_RADIUS = 15

# Regex patterns for sensitive information
REGEX_PATTERNS = {
    "email": r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b",
    "phone_india": r"\b(?:\+?91[-.\\s]?)?(?:0)?[6-9]\d{9}\b",
    "phone_generic": r"\b(?:\+?\d{1,3}[-.\\s]?)?(?:\(?\d{3}\)?[-.\\s]?)?\d{3}[-.\\s]?\d{4}\b",
    "aadhaar": r"\b\d{4}\s?\d{4}\s?\d{4}\b",
    "pan": r"\b[A-Z]{5}[0-9]{4}[A-Z]\b",
    "ssn": r"\b\d{3}-\d{2}-\d{4}\b",
    "credit_card": r"\b(?:\d[ -]*?){13,16}\b",
    "vit_roll": r"\b\d{2}[A-Z]{3}\d{4}\b",
    "generic_roll": r"\b[A-Z]{2,4}\d{3,6}\b",
}

# Entity types to mask
SENSITIVE_ENTITIES = {"PERSON", "ORG", "GPE", "LOC", "DATE", "TIME", "MONEY", "CARDINAL"}


def mask_text(text: str) -> str:
    """
    Mask sensitive information in text using regex and spaCy NER.

    Args:
        text: Input text to sanitize

    Returns:
        Sanitized text with sensitive info replaced with [MASKED]
    """
    if not text or not text.strip():
        return text

    # Apply regex patterns
    masked = text
    for pattern_name, pattern in REGEX_PATTERNS.items():
        masked = re.sub(pattern, "[MASKED]", masked, flags=re.IGNORECASE)

    # Apply spaCy NER
    doc = nlp(masked)
    entities = sorted(doc.ents, key=lambda e: e.start_char, reverse=True)

    for ent in entities:
        if ent.label_ in SENSITIVE_ENTITIES:
            masked = masked[:ent.start_char] + "[MASKED]" + masked[ent.end_char:]

    return masked


def detect_and_mask_faces(image_bytes: bytes) -> bytes:
    """
    Detect and blur faces in an image.

    Args:
        image_bytes: Image data as bytes

    Returns:
        Processed image bytes with faces blurred
    """
    if not MASK_FACES:
        return image_bytes

    try:
        # Convert bytes to numpy array
        arr = np.frombuffer(image_bytes, dtype=np.uint8)
        img = cv2.imdecode(arr, cv2.IMREAD_COLOR)

        if img is None:
            return image_bytes

        # Detect faces
        gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)
        face_cascade = cv2.CascadeClassifier(
            cv2.data.haarcascades + 'haarcascade_frontalface_default.xml'
        )
        faces = face_cascade.detectMultiScale(
            gray, scaleFactor=1.1, minNeighbors=4, minSize=(30, 30)
        )

        # Blur detected faces
        for (x, y, w, h) in faces:
            face_region = img[y:y+h, x:x+w]
            kernel_size = max(15, (w // 7) | 1)
            blurred_face = cv2.GaussianBlur(face_region, (kernel_size, kernel_size), 0)
            img[y:y+h, x:x+w] = blurred_face

        # Encode back to bytes
        success, encoded = cv2.imencode('.jpg', img)
        if success:
            return encoded.tobytes()

    except Exception as e:
        print(f"Warning: Face detection failed: {e}")

    return image_bytes


def sanitize_pptx(input_path: str, output_path: str):
    """
    Sanitize PowerPoint file by masking sensitive information.

    Args:
        input_path: Path to input PPTX file
        output_path: Path to save sanitized PPTX
    """
    prs = Presentation(input_path)

    for slide_num, slide in enumerate(prs.slides, 1):
        shapes_to_process = list(slide.shapes)

        for shape in shapes_to_process:
            # Process text frames
            if hasattr(shape, "has_text_frame") and shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if run.text and run.text.strip():
                            original = run.text
                            masked = mask_text(original)
                            if masked != original:
                                run.text = masked

            # Process tables
            if shape.shape_type == MSO_SHAPE_TYPE.TABLE:
                table = shape.table
                for row_idx in range(len(table.rows)):
                    for col_idx in range(len(table.columns)):
                        cell = table.cell(row_idx, col_idx)
                        if cell.text and cell.text.strip():
                            original = cell.text
                            masked = mask_text(original)
                            if masked != original:
                                cell.text = masked

            # Process images
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    image_bytes = shape.image.blob
                    processed_bytes = detect_and_mask_faces(image_bytes)

                    if processed_bytes != image_bytes:
                        # Save position and size
                        left, top, width, height = shape.left, shape.top, shape.width, shape.height

                        # Remove old shape
                        sp = shape._element
                        sp.getparent().remove(sp)

                        # Add processed image
                        image_stream = io.BytesIO(processed_bytes)
                        slide.shapes.add_picture(image_stream, left, top, width=width, height=height)

                except Exception as e:
                    print(f"Warning: Image processing failed on slide {slide_num}: {e}")

    prs.save(output_path)
    print(f"✓ PPTX sanitized: {output_path}")


if __name__ == "__main__":
    import sys
    if len(sys.argv) != 3:
        print("Usage: python sanitize_pptx.py <input_pptx> <output_pptx>")
        sys.exit(1)

    input_file = sys.argv[1]
    output_file = sys.argv[2]

    print(f"Sanitizing {input_file}...")
    sanitize_pptx(input_file, output_file)